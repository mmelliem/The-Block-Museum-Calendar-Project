import sys
import numpy as np
import pandas as pd
from datetime import datetime
import tkinter as tk
from tkinter import filedialog
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re
import ctypes

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(2)
except:
    pass

root = tk.Tk()
root.withdraw()
root.attributes('-topmost', True)
root.update()
csv_path = filedialog.askopenfilename(title = "Select CSV file", filetypes=[('CSV files', '*.csv'), ('All files', '*.*')])
if not csv_path:
    sys.exit()

df = pd.read_csv(csv_path, encoding='utf-8-sig')

cols = ['Subject', 'Start Date', 'Start Time', 'End Time', 'Description', 'Location']
df = df[cols]
df = df.fillna('')

df['start_dt'] = pd.to_datetime(df['Start Date'] + ' ' + df['Start Time'], format='%m/%d/%Y %I:%M:%S %p')
df['end_dt'] = pd.to_datetime(df['Start Date'] + ' ' + df['End Time'], format='%m/%d/%Y %I:%M:%S %p')
df = df.drop(['Start Date', 'Start Time', 'End Time'], axis = 1)

df = df.sort_values('start_dt').reset_index(drop=True)

df['Subject'] = df['Subject'].str.replace('\n', ' ').str.strip()
mask = ~df['Subject'].str.strip().str.upper().str.startswith(('CANCELED', 'AUTOMATED'))
mask &= ~df['Subject'].str.upper().str.contains('OOO', na=False)
mask &= ~df['Subject'].str.contains('Required Fields Template', case=False, na=False)
df = df[mask]

#df['Description'] = df['Description'].str.replace('\r\n', ' ').str.strip()
df['Description'] = df['Description'].str.strip()
avoid_strings = ['urldefense', 'zoom.com', 'zoom.us', 'teams.microsoft.com']
pattern = '|'.join(avoid_strings)
df['Description'] = df['Description'].mask(
    df['Description'].str.contains(pattern, case=False, na=False), ''
)


req_fields = [
    'Date',
    'Time',
    'Type of Event',
    'Name of Group',
    'Location',
    'Internal Contact Person(s)',
    'External Contact Person(s)',
    'Guests Expected',
    'Taking Attendance (Y/N)',
    'Attendance Total',
    'Notes'
]


def reformat_description(row):
    """Extract values and rebuild with required fields template"""
    try:
        desc = row['Description']
        start_dt = row['start_dt']
        end_dt = row['end_dt']
        
        if not desc or not isinstance(desc, str) or desc.strip() == '':
            return ''
        
        print(f"debug: raw description:\n{repr(desc)}\n")  # script use only

        # Replace weird spaces
        desc = desc.replace('\xa0', ' ')
        
        # Extract values only from lines that start with a known field
        extracted = {}
        current_field = None
        current_value = ''
        
        for line in desc.split('\n'):
            print(f"debug: processing line: {repr(line)}") # script use only
            # Check if line starts with a known field
            field_found = False
            for req_field in req_fields:
                if line.strip().lower().replace(' ', '').startswith((req_field.lower() + ':').replace(' ', '')):
                    # Save previous field if any
                    if current_field:
                        extracted[current_field] = current_value.strip()
                    # Start new field
                    current_field = req_field
                    parts = line.split(':', 1)
                    current_value = parts[1].strip() if len(parts) > 1 else ''
                    field_found = True
                    break
            
            # If no field found, append to current field only if no colon
            if not field_found and current_field:
                if ':' not in line and line.strip():
                    current_value += ' ' + line.strip()
        
        # Save last field
        if current_field:
            extracted[current_field] = current_value.strip()
        
        # If no recognized fields found, return empty
        if not extracted:
            return ''
        
        # Override Date and Time with formatted CSV values
        extracted['Date'] = start_dt.strftime('%m/%d/%Y')
        start_hour = int(start_dt.strftime('%I'))
        end_hour = int(end_dt.strftime('%I'))
        extracted['Time'] = f"{start_hour}:{start_dt.strftime('%M')}{start_dt.strftime('%p').upper()}-{end_hour}:{end_dt.strftime('%M')}{end_dt.strftime('%p').upper()}"
        
        # Build result with all fields
        result = []
        for field in req_fields:
            value = extracted.get(field, '')
            result.append(f"{field}: {value}")
        
        return '\n'.join(result)
    except Exception as e:
        print(f"Error: {e}")
        return ""

    
df['Description'] = df.apply(reformat_description, axis=1)
print("Descriptions updated")
print(df['Description'].head())

def add_horizontal_line(para):
    pPr = para._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "000000")
    pBdr.append(bottom)
    pPr.append(pBdr)

def set_spacing(para, before=0, after=0):
    pPr = para._p.get_or_add_pPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:before"), str(before))
    spacing.set(qn("w:after"), str(after))
    pPr.append(spacing)

def fmt_time(dt):
    return dt.strftime("%I:%M %p").lstrip("0")

doc = Document()

for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.gutter = Inches(0)

style = doc.styles["Normal"]
style.font.name = "Arial"
style.font.size = Pt(12)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.space_after = Pt(0)

p = doc.add_paragraph()
p.add_run().font.size = Pt(12)

# title
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Weekly Events Attendance Totals")
run.font.name = "Arial"
run.font.size = Pt(12)
run.italic = True

doc.add_paragraph() # one blank line

last_date = None

df['date_str'] = df['start_dt'].dt.strftime('%A, %B, %d')
df['time_str'] = (df['start_dt']).dt.strftime('%-I:%M %p') + ' - ' + df['end_dt'].dt.strftime('%-I:%M %p')

for _, row in df.iterrows():
    current_date = row['start_dt'].date()

    if current_date != last_date:
        last_date = current_date
        # date header (size 14, left indent 0.13")
        date_label = row["start_dt"].strftime("%A, %B %d, %Y")
        date_para = doc.add_paragraph()
        date_para.paragraph_format.left_indent = Inches(0.13)
        set_spacing(date_para)
        run = date_para.add_run(date_label)
        run.font.name = "Segoe UI Semibold"
        run.font.size = Pt(14)
        run.bold = True

        # horizontal line
        line_para = doc.add_paragraph()
        line_para.paragraph_format.left_indent = Inches(0.13)
        line_para.paragraph_format.space_before = Pt(0)
        line_para.paragraph_format.space_after = Pt(0)
        line_para.paragraph_format.line_spacing = Pt(1)
        set_spacing(line_para)
        add_horizontal_line(line_para)

    # size-9 event line
    time_str = f"{fmt_time(row['start_dt'])} - {fmt_time(row['end_dt'])}"
    event_para = doc.add_paragraph()
    event_para.paragraph_format.left_indent = Inches(0.13)
    set_spacing(event_para)

    # use location from Outlook, or get from description if empty
    location = str(row['Location']).strip()
    if not location:
        for line in row['Description'].split('\n'):
            if line.startswith('Location:'):
                location = line.split(':', 1)[1].strip()
                break
    run = event_para.add_run(f"{time_str}    {str(row['Subject']).strip()} — {location}")

    run.font.name = "Gadugi"
    run.font.size = Pt(9)
    run.bold = True

    # description lines (italic, "Date" bolded, first line has 7.4pt space before)
    desc_lines = [l for l in str(row["Description"]).split("\n") if l.strip()]
    for i, line in enumerate(desc_lines):
        desc_para = doc.add_paragraph()
        desc_para.paragraph_format.left_indent = Inches(0.31)
        set_spacing(desc_para, before=(148 if i == 0 else 0)) # 7.4pt = 148 twips
        for part in re.split(r'\b(Date)\b', line):
            run = desc_para.add_run(part)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.italic = True
            run.bold = (part == "Date")

    p1 = doc.add_paragraph()
    p1.add_run().font.size = Pt(9)
    p2 = doc.add_paragraph()
    p2.add_run().font.size = Pt(9)


section = doc.sections[0]
sectPr = section._sectPr

for hdr_ref in sectPr.findall(qn("w:headerReference")):
    sectPr.remove(hdr_ref)

# margins
pgMar = sectPr.find(qn("w:pgMar"))
if pgMar is None:
    pgMar = OxmlElement("w:pgMar")
    sectPr.append(pgMar)
pgMar.set(qn("w:top"), "0")
pgMar.set(qn("w:bottom"), str(int(0.19 * 1440))) 
pgMar.set(qn("w:left"), str(int(0.25 * 1440))) 
pgMar.set(qn("w:right"), str(int(1 * 1440)))
pgMar.set(qn("w:header"), "0")
pgMar.set(qn("w:footer"), "0")
pgMar.set(qn("w:gutter"), "0")


output_path = filedialog.asksaveasfilename(title="Save Word document as", defaultextension='.docx', filetypes=[('Word files','*.docx'), ('All files', '*.*')])
if not output_path:
    sys.exit()
doc.save(output_path)
root.destroy()